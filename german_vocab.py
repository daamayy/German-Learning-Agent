"""
German Vocabulary Extractor
Automatically extracts German vocabulary from YouTube videos and saves to Excel.

Dependencies:
    pip install youtube-transcript-api openpyxl google-genai
"""

# ── Gemini API key ────────────────────────────────────────────────────────────
# No need to edit this. The key is stored in config.json in the same folder
GEMINI_API_KEY = ""
# ─────────────────────────────────────────────────────────────────────────────

import tkinter as tk
from tkinter import ttk, messagebox, scrolledtext
import threading
import json
import os
import re
from datetime import datetime
from pathlib import Path


def check_and_import():
    missing = []
    try:
        import google.genai as _
    except ImportError:
        missing.append("google-genai")
    try:
        import youtube_transcript_api
    except ImportError:
        missing.append("youtube-transcript-api")
    try:
        import openpyxl
    except ImportError:
        missing.append("openpyxl")
    return missing


# ─────────────────────────────────────────────
# Core Logic
# ─────────────────────────────────────────────

def extract_video_id(url: str) -> str | None:
    patterns = [r"(?:v=|youtu\.be/|embed/)([A-Za-z0-9_-]{11})"]
    for p in patterns:
        m = re.search(p, url)
        if m:
            return m.group(1)
    return None


def get_transcript(video_id: str) -> tuple[str, str]:
    """Returns (transcript_text, language_used)"""
    from youtube_transcript_api import YouTubeTranscriptApi

    api = YouTubeTranscriptApi()

    # 优先德语，其次英语，最后自动
    for lang in [["de"], ["en"], None]:
        try:
            if lang:
                entries = api.fetch(video_id, languages=lang)
            else:
                transcript_list = api.list(video_id)
                entries = transcript_list.find_transcript(
                    transcript_list._manually_created_transcripts or
                    transcript_list._generated_transcripts
                ).fetch()
            text = " ".join(e.text for e in entries)
            used = lang[0] if lang else "auto"
            return text, used
        except Exception:
            continue

    raise RuntimeError("Could not retrieve subtitles. Please ensure the video has public captions.")


SYSTEM_PROMPT = """⚠️ MANDATORY COMPLIANCE — READ THIS FIRST, BEFORE DOING ANYTHING ELSE:
You MUST output between 50 and 80 complete JSON entries. This is a hard requirement.
- DO NOT stop after a few entries. Stopping before 50 entries is total task failure.
- DO NOT abbreviate, summarize, or skip entries due to perceived length constraints.
- CONTINUE scanning and generating until you have reached the end of the transcript.
- Every single entry MUST have every field fully filled out. No shortcuts.
- You will be penalized for incomplete output. There is no penalty for being thorough.

You are an expert German language teacher and computational linguist. Your task is to analyze a YouTube video transcript provided by the user, perform a full-text scan, and extract high-value German vocabulary, idioms, and particles to generate structured learning data.

### CRITICAL BEHAVIOR ADJUSTMENTS (Read Carefully)
FULL-TEXT SCANNING: Do not just focus on the beginning of the transcript. You must scan the entire text from start to finish. Ensure the extracted words are evenly distributed across the entire timeline of the video.

### JSON OUTPUT FORMAT
Respond strictly with a single JSON object. Do not include any markdown formatting (such as ```json), wrap-up text, or explanations. 

{
  "video_title": "A concise Deutsch summary of the video theme based on the transcript (under 10 characters).",
  "entries": [
    {
      "vocabulary": "ankommen auf",
      "verb_conjugation": "kommt an / kam an / ist angekommen",
      "part_of_speech": "v.",
      "english_meaning": "to depend on, to matter",
      "original_context_sentence": "Es kommt darauf an, wie du es machst.",
      "original_sentence_translation": "It depends on how you do it.",
      "particle_or_collocation": "Es kommt auf etw. (Akk.) an / darauf ankommen",
      "cefr_level": "B1",
      "grammar_notes": "可分动词，常与代副词 darauf 搭配"
    }
  ]
}

### MUST-EXTRACT CATEGORIES (Prioritize these)
1. Separable verbs with fixed prepositional collocations (e.g., ankommen auf, ausgehen von, umgehen mit, sich beziehen auf).
2. Pronominal Adverbs (Pronominaladverbien): daran, davon, damit, dafür, darauf, darüber, dazu, danach, etc. (Explain their grammatical function in 'grammar_notes').
3. Modal Particles (Modalpartikeln): doch, mal, ja, eben, halt, eigentlich, schon, bloß, etc. (Crucial! Explain the exact nuanced tone and conversational context in 'grammar_notes').
4. High-frequency prepositional phrases (e.g., im Vergleich zu, aufgrund von, im Hinblick auf).
5. Essential vocabulary from CEFR levels A2 to C1.
6. High-frequency colloquial/spoken words that might not have a strict CEFR level (mark level as "frequently used").
7. Any word above A1 level that a German learner would benefit from knowing — even if it appears only once in the transcript. Do NOT skip a word simply because it is uncommon or occurs infrequently; rarity in the transcript is not a reason to exclude it.
8. Essential conjunctions (e.g., obwohl, während, sobald) and adverbs (e.g., trotzdem, außerdem, allerdings) that control sentence structure.

### DO NOT EXTRACT
- Ultra-basic A1 words (e.g., sein, haben, gehen, kommen, ich, du, und, aber, nicht, auch).
- Pure English words.

### EXTRACTION RULES
- Target Quantity: Scan the entire transcript thoroughly and extract 50 to 80 high-quality entries. Do not stop until you have covered the end of the video.
- Selection Criterion: Select words based on their EDUCATIONAL VALUE for a German learner, not their frequency in this transcript. A word that appears only once but is above A1 level MUST be included if a learner would benefit from knowing it.
- Nouns must include their definite article (e.g., der Bahnhof). Verbs must be in the infinitive form.
- The 'original_context_sentence' must be a complete, meaningful sentence from the transcript, not just a fragmented phrase.
- CEFR Levels allowed: A2 / B1 / B2 / C1 / frequently used (Never use A1).
- Part of Speech Abbreviations: n.=noun, v.=verb, adj.=adjective, adv.=adverb, prep.=preposition, conj.=conjunction, part.=particle, pron.adv.=pronominal adverb, pron.=pronoun.
- 'verb_conjugation': ONLY fill this field when the entry is a verb (v.). Format: "er-form present / ich-form Präteritum / Perfekt auxiliary + Partizip II". Example: "kommt an / kam an / ist angekommen". Leave as empty string "" for all non-verb entries.
- 'original_sentence_translation': Write the translation in ENGLISH.
- 'grammar_notes': Write ALL grammar notes in CHINESE (中文)."""


def _safe_json_parse(raw: str) -> dict:
    """Parse JSON from Gemini; fix common quirks and recover truncated output."""
    # Gemini 偶尔输出 Python 字面量而非 JSON 字面量
    raw = re.sub(r'\bNone\b',  'null',  raw)
    raw = re.sub(r'\bTrue\b',  'true',  raw)
    raw = re.sub(r'\bFalse\b', 'false', raw)

    try:
        return json.loads(raw)
    except json.JSONDecodeError as exc:
        # 尝试从截断的 JSON 中恢复完整条目
        last = raw.rfind("\n    }")
        if last != -1:
            try:
                return json.loads(raw[: last + 6] + "\n  ]\n}")
            except json.JSONDecodeError:
                pass
        # 报错时附上出错位置附近的原文，便于调试
        lo  = max(0, exc.pos - 120)
        hi  = min(len(raw), exc.pos + 120)
        ctx = raw[lo:hi].replace("\n", "↵")
        raise ValueError(
            f"JSON parse error — {exc.msg} at char {exc.pos}\n"
            f"Context: ...{ctx}..."
        ) from exc


def extract_vocab_with_gemini(transcript: str, api_key: str, video_url: str) -> list[dict]:
    import time
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=api_key)

    # gemini-2.5-flash 上下文窗口 100万 token，字幕全量发送即可；
    # 超长视频（>50K字符）才截断，避免极端情况
    if len(transcript) > 50_000:
        transcript = transcript[:50_000] + "...[truncated]"

    # 遇到 503/429 等临时错误自动重试，最多 3 次，指数退避
    last_err = None
    for attempt in range(3):
        try:
            response = client.models.generate_content(
                model="gemini-2.5-flash",
                config=types.GenerateContentConfig(
                    system_instruction=SYSTEM_PROMPT,
                    max_output_tokens=20000,
                    response_mime_type="application/json",  # 强制约束解码，防止早停和格式错误
                ),
                contents=f"Please analyze the following German video transcript:\n\n{transcript}",
            )
            break  # 成功则跳出重试循环
        except Exception as e:
            last_err = e
            msg = str(e)
            # 只对临时性错误重试（503 过载、429 限速）
            if "503" in msg or "429" in msg or "UNAVAILABLE" in msg or "RESOURCE_EXHAUSTED" in msg:
                time.sleep(10 * (2 ** attempt))  # 10s → 20s → 40s
                continue
            raise  # 其他错误（401、404 等）直接抛出，不重试
    else:
        raise last_err  # 3 次全失败

    raw = response.text.strip()

    raw = re.sub(r"^```(?:json)?\s*", "", raw)
    raw = re.sub(r"\s*```$", "", raw)

    data = _safe_json_parse(raw)
    entries = data.get("entries", [])
    title_guess = data.get("video_title", data.get("video_title_guess", "Unknown"))

    now = datetime.now().strftime("%Y-%m-%d")
    for e in entries:
        e["video_url"] = video_url
        e["topic"] = title_guess
        e["date_added"] = now

    return entries


# (Gemini JSON key, Excel column header)
_COLUMNS = [
    ("vocabulary",                    "Word"),
    ("verb_conjugation",              "Conjugation"),
    ("part_of_speech",                "Part of Speech"),
    ("english_meaning",               "English Meaning"),
    ("original_context_sentence",     "Context Sentence"),
    ("original_sentence_translation", "Context Translation"),
    ("particle_or_collocation",       "Particles / Collocations"),
    ("cefr_level",                    "CEFR Level"),
    ("grammar_notes",                 "Grammar Notes"),
    ("topic",                         "Topic"),
    ("video_url",                     "Video URL"),
    ("date_added",                    "Date Added"),
]

_COL_WIDTHS   = [22, 28, 12, 22, 40, 22, 35, 10, 35, 18, 32, 12]
_LEFT_ALIGNED = {
    "original_context_sentence", "original_sentence_translation",
    "particle_or_collocation", "grammar_notes", "video_url",
}


def append_to_excel(entries: list[dict], filepath: str) -> tuple[int, int]:
    """Append entries to Excel, skipping duplicates. Returns (added, skipped)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    header_fill = PatternFill("solid", start_color="7C83D3")
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    thin   = Side(style="thin", color="E5E7EB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    existing_words: set[str] = set()

    if os.path.exists(filepath):
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        # 读取已有单词用于去重（跳过表头）
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0]:
                existing_words.add(str(row[0]).strip().lower())
        next_row = ws.max_row + 1
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "German Vocabulary"
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "A2"

        for col_idx, (_, header) in enumerate(_COLUMNS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border
            ws.column_dimensions[cell.column_letter].width = _COL_WIDTHS[col_idx - 1]

        next_row = 2

    fill_light = PatternFill("solid", start_color="EEF2FF")
    fill_white = PatternFill("solid", start_color="FFFFFF")
    added = skipped = 0

    for entry in entries:
        word = str(entry.get("vocabulary", "")).strip().lower()
        if word and word in existing_words:
            skipped += 1
            continue

        row_fill = fill_light if (next_row % 2 == 0) else fill_white
        for col_idx, (key, _) in enumerate(_COLUMNS, 1):
            val = entry.get(key, "")
            cell = ws.cell(row=next_row, column=col_idx, value=val)
            cell.border = border
            cell.fill = row_fill
            cell.font = Font(name="Arial", size=10)
            cell.alignment = left if key in _LEFT_ALIGNED else center

        if word:
            existing_words.add(word)
        next_row += 1
        added += 1

    wb.save(filepath)
    return added, skipped


# ─────────────────────────────────────────────
# GUI — Minimalist Pastel Design
# ─────────────────────────────────────────────

C = {
    "bg":             "#F7F8FC",
    "card":           "#FFFFFF",
    "header_bg":      "#EEF2FF",
    "accent":         "#5C63B8",   # deeper indigo for contrast
    "accent_hover":   "#4A51A0",
    "mint":           "#A8D5BA",
    "peach":          "#F5C6A0",
    "lavender":       "#C5C8F0",
    "text":           "#1A1D2E",   # near-black
    "text_mid":       "#374151",   # dark gray
    "text_light":     "#6B7280",   # medium gray (was too light)
    "border":         "#D1D5DB",
    "log_bg":         "#1A1D2E",
    "log_fg":         "#A5B4FC",
}

FF = "Helvetica Neue"   # macOS; falls back gracefully


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("German Vocabulary Extractor")
        self.geometry("760x660")
        self.minsize(680, 560)
        self.configure(bg=C["bg"])
        self.resizable(True, True)

        # 配置就放在本脚本所在的文件夹里，文件名 config.json（可见，不隐藏）
        self.config_file = Path(__file__).resolve().parent / "config.json"
        self._migrate_old_config()
        self.config = self._load_config()

        self._setup_styles()
        self._build_ui()
        self._check_deps()

    def _migrate_old_config(self):
        """把旧的隐藏配置 ~/.german_vocab_config.json 搬到脚本目录下的 config.json"""
        old = Path.home() / ".german_vocab_config.json"
        if old.exists() and not self.config_file.exists():
            try:
                self.config_file.write_text(old.read_text())
            except Exception:
                pass

    def _load_config(self) -> dict:
        if self.config_file.exists():
            try:
                return json.loads(self.config_file.read_text())
            except Exception:
                pass
        # 第一次运行：自动生成一份空的 config.json，方便直接打开填写
        cfg = {"api_key": "", "output_path": str(Path.home() / "Desktop" / "German_Vocabulary.xlsx")}
        try:
            self.config_file.write_text(json.dumps(cfg, ensure_ascii=False, indent=2))
        except Exception:
            pass
        return cfg

    def _save_config(self):
        self.config["output_path"] = self.output_path_var.get().strip()
        self.config_file.write_text(json.dumps(self.config, ensure_ascii=False, indent=2))

    def _setup_styles(self):
        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure(
            "Pastel.Horizontal.TProgressbar",
            troughcolor=C["border"],
            background=C["lavender"],
            bordercolor=C["border"],
            lightcolor=C["lavender"],
            darkcolor=C["lavender"],
        )

    # ── Layout helpers ───────────────────────

    def _card(self, parent, label: str) -> tk.Frame:
        """Returns an inner content frame inside a white card with a section label."""
        outer = tk.Frame(parent, bg=C["card"])
        outer.pack(fill="x", pady=(0, 8))

        label_row = tk.Frame(outer, bg=C["card"], padx=16, pady=10)
        label_row.pack(fill="x")
        tk.Label(
            label_row,
            text=label.upper(),
            font=(FF, 9, "bold"),
            fg=C["text_mid"],
            bg=C["card"],
        ).pack(anchor="w")
        tk.Frame(outer, bg=C["border"], height=1).pack(fill="x", padx=16)

        inner = tk.Frame(outer, bg=C["card"], padx=16, pady=12)
        inner.pack(fill="x")
        return inner

    def _label(self, parent, text, width=14) -> tk.Label:
        return tk.Label(
            parent, text=text, width=width, anchor="w",
            font=(FF, 11), fg=C["text"], bg=C["card"],
        )

    def _entry(self, parent, textvariable, show=None, width=44) -> tk.Entry:
        e = tk.Entry(
            parent, textvariable=textvariable,
            font=(FF, 11), width=width,
            bg="#F3F4F6", fg=C["text"],
            relief="flat", bd=0,
            highlightthickness=1,
            highlightbackground=C["border"],
            highlightcolor=C["accent"],
            insertbackground=C["accent"],
        )
        if show:
            e.config(show=show)
        return e

    def _button(self, parent, text, command, primary=False) -> tk.Button:
        return tk.Button(
            parent, text=text,
            font=(FF, 10, "bold") if primary else (FF, 10),
            bg=C["accent"] if primary else C["border"],
            fg="white" if primary else C["text_mid"],
            activebackground=C["accent_hover"] if primary else "#C1C5CB",
            activeforeground="white" if primary else C["text"],
            relief="flat", bd=0, cursor="hand2",
            padx=16, pady=6,
            command=command,
        )

    # ── UI construction ──────────────────────

    def _build_ui(self):
        # Header
        header = tk.Frame(self, bg=C["header_bg"])
        header.pack(fill="x")
        inner = tk.Frame(header, bg=C["header_bg"], padx=24, pady=18)
        inner.pack(fill="x")
        tk.Label(
            inner, text="German Vocabulary Extractor",
            font=(FF, 19, "bold"), fg=C["accent"], bg=C["header_bg"],
        ).pack(anchor="w")
        tk.Label(
            inner,
            text="Paste a YouTube link  ·  AI extracts vocabulary  ·  Saved to Excel",
            font=(FF, 11), fg=C["text_mid"], bg=C["header_bg"],
        ).pack(anchor="w", pady=(4, 0))

        # Accent divider
        tk.Frame(self, bg=C["lavender"], height=2).pack(fill="x")

        # Main content
        content = tk.Frame(self, bg=C["bg"])
        content.pack(fill="both", expand=True, padx=20, pady=14)

        self._build_settings(content)
        self._build_input(content)
        self._build_status(content)
        self._build_log(content)

    def _build_settings(self, parent):
        inner = self._card(parent, "Settings")

        row = tk.Frame(inner, bg=C["card"])
        row.pack(fill="x")
        self._label(row, "Excel Output").pack(side="left")
        self.output_path_var = tk.StringVar(value=self.config.get("output_path", ""))
        self._entry(row, self.output_path_var).pack(side="left", padx=(0, 10))
        self._button(row, "Browse", self._browse_path).pack(side="left")

    def _build_input(self, parent):
        inner = self._card(parent, "YouTube URL")

        row = tk.Frame(inner, bg=C["card"])
        row.pack(fill="x")

        self.url_var = tk.StringVar()
        url_entry = self._entry(row, self.url_var, width=52)
        url_entry.pack(side="left", fill="x", expand=True, padx=(0, 10))
        url_entry.bind("<Return>", lambda e: self._run())

        # Placeholder behaviour
        placeholder = "https://www.youtube.com/watch?v=..."
        url_entry.insert(0, placeholder)
        url_entry.config(fg=C["text_light"])

        def on_focus_in(e):
            if url_entry.get() == placeholder:
                url_entry.delete(0, "end")
                url_entry.config(fg=C["text"])

        def on_focus_out(e):
            if not url_entry.get():
                url_entry.insert(0, placeholder)
                url_entry.config(fg=C["text_light"])

        url_entry.bind("<FocusIn>", on_focus_in)
        url_entry.bind("<FocusOut>", on_focus_out)

        self.run_btn = self._button(row, "▶  Extract", self._run, primary=True)
        self.run_btn.pack(side="left")

    def _build_status(self, parent):
        frame = tk.Frame(parent, bg=C["bg"])
        frame.pack(fill="x", pady=(0, 8))

        self.progress_bar = ttk.Progressbar(
            frame, mode="indeterminate", style="Pastel.Horizontal.TProgressbar"
        )
        self.progress_bar.pack(fill="x", pady=(0, 4))

        self.progress_var = tk.StringVar(value="Ready")
        tk.Label(
            frame, textvariable=self.progress_var,
            font=(FF, 10), fg=C["text_mid"], bg=C["bg"],
        ).pack(anchor="w")

    def _build_log(self, parent):
        outer = tk.Frame(parent, bg=C["card"])
        outer.pack(fill="both", expand=True)

        label_row = tk.Frame(outer, bg=C["card"], padx=16, pady=10)
        label_row.pack(fill="x")
        tk.Label(
            label_row, text="ACTIVITY LOG",
            font=(FF, 9, "bold"), fg=C["text_mid"], bg=C["card"],
        ).pack(anchor="w")
        tk.Frame(outer, bg=C["border"], height=1).pack(fill="x", padx=16)

        log_inner = tk.Frame(outer, bg=C["card"], padx=16, pady=12)
        log_inner.pack(fill="both", expand=True)

        self.log = scrolledtext.ScrolledText(
            log_inner,
            height=10,
            font=("Menlo", 11),
            bg=C["log_bg"],
            fg=C["log_fg"],
            insertbackground="white",
            state="disabled",
            relief="flat",
            bd=0,
            padx=12, pady=10,
        )
        self.log.pack(fill="both", expand=True)

    # ── Actions ──────────────────────────────

    def _browse_path(self):
        from tkinter import filedialog
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="German_Vocabulary.xlsx",
        )
        if path:
            self.output_path_var.set(path)

    def _check_deps(self):
        missing = check_and_import()
        if missing:
            self.log_msg(f"Missing packages — run: pip install {' '.join(missing)}")
        self.log_msg("Ready  ·  Paste a YouTube link and click Extract")
        if not (GEMINI_API_KEY or self.config.get("api_key")):
            self.log_msg(f"⚠  没有找到 API Key —— 请打开 {self.config_file} 填入 api_key")

    def log_msg(self, msg: str):
        self.log.configure(state="normal")
        self.log.insert("end", f"[{datetime.now().strftime('%H:%M:%S')}]  {msg}\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def _run(self):
        placeholder = "https://www.youtube.com/watch?v=..."
        url = self.url_var.get().strip()
        if url == placeholder:
            url = ""

        api_key = GEMINI_API_KEY or self.config.get("api_key", "")
        output_path = self.output_path_var.get().strip()

        if not url:
            messagebox.showwarning("Missing URL", "Please enter a YouTube URL.")
            return
        if not api_key:
            messagebox.showwarning(
                "缺少 API Key",
                f"请用文本编辑器打开下面这个文件，把 Gemini API Key 填进 api_key：\n\n{self.config_file}",
            )
            return
        if not output_path:
            messagebox.showwarning("Missing Path", "Please set an Excel output path.")
            return

        self._save_config()
        self.run_btn.configure(state="disabled", bg=C["text_light"])
        self.progress_bar.start(10)

        threading.Thread(
            target=self._worker, args=(url, api_key, output_path), daemon=True
        ).start()

    def _worker(self, url: str, api_key: str, output_path: str):
        try:
            self._set_status("Parsing video ID...")
            video_id = extract_video_id(url)
            if not video_id:
                raise ValueError("Could not parse video ID — check the URL format.")
            self.log_msg(f"Video ID: {video_id}")

            self._set_status("Fetching transcript...")
            transcript, lang = get_transcript(video_id)
            self.log_msg(f"Language: {lang}  ·  {len(transcript):,} characters")
            self.log_msg(f"Preview: {transcript[:100]}...")

            self._set_status("Analyzing with Gemini AI  (10–20 sec)...")
            entries = extract_vocab_with_gemini(transcript, api_key, url)
            self.log_msg(f"Extracted {len(entries)} vocabulary entries")

            self._set_status("Saving to Excel...")
            added, skipped = append_to_excel(entries, output_path)
            self.log_msg(f"Saved → {output_path}")
            if skipped:
                self.log_msg(f"Skipped {skipped} duplicate(s) already in Excel")

            self.log_msg("─" * 50)
            self.log_msg("Vocabulary preview:")
            for e in entries[:8]:
                self.log_msg(
                    f"  {e.get('vocabulary',''):20s}  {e.get('part_of_speech',''):6s}  {e.get('english_meaning','')}"
                )
            if len(entries) > 8:
                self.log_msg(f"  … and {len(entries) - 8} more")
            self.log_msg("─" * 50)

            if added == 0:
                self._set_status(f"Done  ·  All {skipped} words already in Excel")
            elif skipped:
                self._set_status(f"Done  ·  {added} added, {skipped} duplicates skipped")
            else:
                self._set_status(f"Done  ·  {added} words added to Excel")
            self.after(0, lambda: self.url_var.set(""))

        except Exception as e:
            self.log_msg(f"Error: {e}")
            self._set_status(f"Failed: {str(e)[:70]}")
        finally:
            self.after(0, lambda: self.run_btn.configure(state="normal", bg=C["accent"]))
            self.after(0, self.progress_bar.stop)

    def _set_status(self, msg: str):
        self.log_msg(msg)
        self.after(0, lambda: self.progress_var.set(msg))


if __name__ == "__main__":
    app = App()
    app.mainloop()
